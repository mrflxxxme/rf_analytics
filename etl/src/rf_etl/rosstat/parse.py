"""Разбор выгрузок Росстата для `region_profile` (03-methodology.md, п. 12).

showdata.rosstat.gov.ru — JSON отчёта (`/view/compound/`): строки с кодом ОКАТО, колонки — периоды.
xlsx Росстата — строки с названиями субъектов; сопоставление с `region_id` — `RegionNames`.
Архангельская и Тюменская области берутся без автономных округов: полигоны АО на карте отдельные.
"""

from __future__ import annotations

import csv
import json
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

# region_id → ОКАТО территории без АО (в regions.csv ОКАТО субъекта целиком).
WITHOUT_AO = {"29": "11001", "72": "71001"}


@dataclass(frozen=True)
class Obs:
    region_id: str
    indicator: str
    year: int
    value: float
    kind: str  # семейство источника: showdata | polvozr | okpopul | urov (приоритет — profile.INDICATORS)
    source: str  # файл или отчёт
    computed: bool = False  # посчитано из опубликованных величин (доля, коэффициент), а не взято готовым
    provisional: bool = False  # Росстат помечает как предварительные данные


@dataclass
class Parsed:
    obs: list[Obs] = field(default_factory=list)
    unmatched: list[str] = field(default_factory=list)  # названия строк без субъекта (ФО, РФ, сноски)
    conflicts: int = 0  # ячейки showdata с несколькими значениями (db_conflict) — пропущены


def load_regions(path: str | Path) -> list[dict[str, str]]:
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def okato_index(regions: Iterable[dict[str, str]]) -> dict[str, str]:
    """11-значный код ОКАТО строки showdata → region_id."""
    return {WITHOUT_AO.get(r["region_id"], r["okato"]).ljust(11, "0"): r["region_id"] for r in regions}


_FOOTNOTE = re.compile(r"\d+\)|\*")
_PAREN = re.compile(r"\([^)]*\)")
_CITY = re.compile(r"(^|\s)г\.\s*")
_PUNCT = re.compile(r"[-–—.,]")


def norm_name(s: str) -> str:
    s = s.lower().replace("ё", "е").replace("p", "р")  # латинская p в «Федеpация»
    s = _CITY.sub(" ", _PAREN.sub(" ", _FOOTNOTE.sub(" ", s)))
    s = _PUNCT.sub(" ", s.replace("авт.", "автономный"))
    return " ".join(s.split())


class RegionNames:
    """Название субъекта из xlsx → region_id. Короткая форма («Кемеровская область») — по префиксу."""

    def __init__(self, regions: Iterable[dict[str, str]]):
        self._keys = {norm_name(r["name"]): r["region_id"] for r in regions}

    def resolve(self, name: str) -> tuple[str | None, bool]:
        """(region_id, строка «без автономного округа»)."""
        base, sep, _ = norm_name(name).partition(" без ")
        return self._match(base), bool(sep)

    def _match(self, k: str) -> str | None:
        if k in self._keys:
            return self._keys[k]
        hits = {rid for ck, rid in self._keys.items() if ck.startswith(k + " ") or k.startswith(ck + " ")}
        return hits.pop() if len(hits) == 1 else None


def by_region(names: RegionNames, rows: Iterable[tuple[str, Any]], res: Parsed) -> dict[str, Any]:
    """Строки таблицы → {region_id: payload}. Если есть строка «… без АО», общая строка субъекта отбрасывается."""
    hits: list[tuple[str, bool, Any]] = []
    for name, payload in rows:
        rid, without_ao = names.resolve(name)
        if rid is None:
            res.unmatched.append(name.strip())
        else:
            hits.append((rid, without_ao, payload))
    narrowed = {rid for rid, wo, _ in hits if wo}
    out: dict[str, Any] = {}
    for rid, wo, payload in hits:
        if rid in narrowed and not wo:
            continue
        if rid in out:
            raise ValueError(f"субъект {rid} встречается в таблице дважды")
        out[rid] = payload
    return out


def _num(v: object) -> float | None:
    return float(v) if isinstance(v, int | float) and not isinstance(v, bool) else None


def _name_rows(ws: Any, cols: tuple[int, ...]) -> Iterator[tuple[str, tuple[float | None, ...]]]:
    """Строки листа, где в первой колонке текст, а в колонках `cols` есть хотя бы одно число."""
    for row in ws.iter_rows(values_only=True):
        if not row or not isinstance(row[0], str):
            continue
        vals = tuple(_num(row[c]) if c < len(row) else None for c in cols)
        if any(v is not None for v in vals):
            yield row[0], vals


def _sheet(wb: Any, name: str, index: int) -> Any:
    return wb[name] if name in wb.sheetnames else wb.worksheets[index]


# --- showdata --------------------------------------------------------------------------------

def _leaves(nodes: list[dict[str, Any]]) -> Iterator[dict[str, Any]]:
    for n in nodes:
        if n.get("children"):
            yield from _leaves(n["children"])
        else:
            yield n


def parse_showdata(path: str | Path, indicator: str, okato: dict[str, str]) -> Parsed:
    d = json.loads(Path(path).read_text("utf-8"))
    hdr = d["headers"]["reportHeaders"]
    rows, cols = list(_leaves(hdr["row_header"])), list(_leaves(hdr["col_header"]))
    data = d["data"]["reportData"]["data"]
    if not rows or len(data) != len(rows):
        raise ValueError(f"{path}: строк {len(rows)}, данных {len(data)}")
    years = [int(c["period"][:4]) for c in cols]
    res, src = Parsed(), f"showdata/{Path(path).stem.removeprefix('showdata_')}"
    for row, cells in zip(rows, data, strict=True):
        rid = okato.get(str(row.get("extra_row_attrs", {}).get("code")))
        if rid is None:
            res.unmatched.append(str(row.get("display_title", "")))
            continue
        for year, cell in zip(years, cells, strict=True):
            if "db_conflict" in cell:
                res.conflicts += 1
            elif (v := _num(cell.get("db_value"))) is not None:
                res.obs.append(Obs(rid, indicator, year, v, "showdata", src))
    return res


# --- xlsx: численность по полу и возрасту (пересчёт после ВПН-2020) --------------------------

POLVOZR = re.compile(r"Chisl_polvozr_01-01-(\d{4})_VPN-2020\.xlsx$")


def _urban_share(rid: str, year: int, urban: float | None, total: float | None, kind: str, src: str) -> list[Obs]:
    if urban is None or not total:
        return []
    return [Obs(rid, "urban_share", year, round(100 * urban / total, 2), kind, src, computed=True)]


def parse_polvozr(path: str | Path, names: RegionNames) -> Parsed:
    """Численность на 1 января года из имени файла: лист М+Ж_Г+С (всё), М+Ж_Г (городское)."""
    p = Path(path)
    m = POLVOZR.search(p.name)
    if not m:
        raise ValueError(f"не файл численности по полу и возрасту: {p.name}")
    year, res = int(m.group(1)), Parsed()
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    total = by_region(names, _name_rows(_sheet(wb, "М+Ж_Г+С", 0), (1,)), res)
    urban = by_region(names, _name_rows(_sheet(wb, "М+Ж_Г", 3), (1,)), Parsed())
    for rid, (t,) in total.items():
        if t is not None:
            res.obs.append(Obs(rid, "population", year, t, "polvozr", p.name))
        res.obs += _urban_share(rid, year, (urban.get(rid) or (None,))[0], t, "polvozr", p.name)
    return res


# --- xlsx: оценка численности и компоненты изменения (OkPopul_CompYYYY_Site) ------------------

OKPOPUL = re.compile(r"OkPopul_Comp(\d{4})_Site\.xlsx$")
# Колонки: 1 — на 1 января Y−1, 4 — миграционный прирост за Y−1, 6 — на 1 января Y, 7 — в среднем за Y−1.
_OK_COLS = (1, 4, 6, 7)


def parse_okpopul(path: str | Path, names: RegionNames) -> Parsed:
    p = Path(path)
    m = OKPOPUL.search(p.name)
    if not m:
        raise ValueError(f"не файл оценки численности: {p.name}")
    year, res = int(m.group(1)), Parsed()
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    total = by_region(names, _name_rows(_sheet(wb, "Всего", 0), _OK_COLS), res)
    urban = by_region(names, _name_rows(_sheet(wb, "Город", 1), _OK_COLS), Parsed())
    for rid, (p0, mig, p1, avg) in total.items():
        u0, _, u1, _ = urban.get(rid, (None,) * 4)
        for y, t, u in ((year - 1, p0, u0), (year, p1, u1)):
            if t is not None:
                res.obs.append(Obs(rid, "population", y, t, "okpopul", p.name))
            res.obs += _urban_share(rid, y, u, t, "okpopul", p.name)
        if mig is not None and avg:
            res.obs.append(Obs(rid, "migration_rate", year - 1, round(1e4 * mig / avg, 1), "okpopul",
                               p.name, computed=True))
    return res


# --- xlsx: среднедушевые денежные доходы по субъектам (urov_10subg_*) ------------------------

INCOME = re.compile(r"urov_10subg_[\w-]+\.xlsx$", re.I)
_YEAR = re.compile(r"(20\d\d)\s*год")


def _annual_columns(rows: list[tuple[Any, ...]]) -> tuple[int, dict[int, tuple[int, bool]]]:
    """Индекс строки с подзаголовками и {колонка: (год, предварительные)} для колонок «год»."""
    for i, row in enumerate(rows[:-1]):
        heads = ["" if c is None else str(c) for c in row]
        if not any(_YEAR.search(h) for h in heads):
            continue
        sub = ["" if c is None else str(c).strip() for c in rows[i + 1]]
        cols: dict[int, tuple[int, bool]] = {}
        year, prov = None, False
        for j, h in enumerate(heads):
            if (ym := _YEAR.search(h)) is not None:
                year, prov = int(ym.group(1)), bool(_FOOTNOTE.search(h))
            label = sub[j] if j < len(sub) else ""
            if year is not None and (label == "год" or (not label and ym is not None)):
                cols[j] = (year, prov)
        return i + 1, cols
    raise ValueError("не нашли строку с годами")


def parse_income(path: str | Path, names: RegionNames) -> Parsed:
    """Лист СДД_субъекты: руб. в месяц, колонки «год» (кварталы и полугодия не берём)."""
    p, res = Path(path), Parsed()
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = _sheet(wb, "СДД_субъекты", 1 if len(wb.sheetnames) > 1 else 0)
    rows = list(ws.iter_rows(values_only=True))
    sub_i, cols = _annual_columns(rows)
    picked = by_region(names, ((r[0], r) for r in rows[sub_i + 1:] if r and isinstance(r[0], str)
                               and any(_num(v) is not None for v in r[1:])), res)
    for rid, row in picked.items():
        for j, (year, prov) in cols.items():
            if j < len(row) and (v := _num(row[j])) is not None:
                res.obs.append(Obs(rid, "income_per_capita", year, v, "urov", p.name, provisional=prov))
    return res
